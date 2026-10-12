"""Build a one-file CV as .docx, in English or German, with no LaTeX install.

The name is historical: it was written to merge two designed CVs (one per role
class) into a single combined one. What it does is generic - render the section
list in DOCX_CV[lang] from application/cv_data.py (see cv_data.example.py) under a
header built from NAME / CONTACT / ADDRESS / LINKS in the same file.

Nothing here is invented: take the wording from your designed CVs and from
application/resume.md.

    PYTHONUTF8=1 python scripts/make_combined_cv.py [--lang en|de] [<out.docx>]

then convert with scripts/docx2pdfa.ps1. Default output is CV_<FILE_STEM>.docx
(en) or Lebenslauf_<FILE_STEM>.docx (de) in the current folder.

Each section entry is (kind, *args), dispatched to the helpers in build():
head/bullet/text take one string, entry and skill take two.
"""
import argparse
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_tailored_cv import load_data, stem, strip_marks  # noqa: E402

ACCENT = RGBColor(0x1F, 0x3B, 0x63)
GREY = RGBColor(0x44, 0x44, 0x44)
SEP = '  |  '


def hr(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), '6')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '1F3B63')
    b.append(bot)
    pPr.append(b)


def build(out, d, lang='en'):
    c = d.DOCX_CV[lang]
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(9.5)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.1)
        s.left_margin = s.right_margin = Cm(1.5)

    def para(space_after=2, space_before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = 1.0
        return p

    def head(text):
        p = para(space_after=4, space_before=8)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = ACCENT
        hr(p)

    def entry(title, meta=None):
        p = para(space_after=1, space_before=4)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10)
        if meta:
            r2 = p.add_run('   ' + meta)
            r2.italic = True
            r2.font.size = Pt(9)
            r2.font.color.rgb = GREY

    def bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.font.size = Pt(9.5)
        return p

    def skill(label, items):
        p = para(space_after=2)
        r = p.add_run(label + ' — ')
        r.bold = True
        r.font.size = Pt(9.5)
        r2 = p.add_run(items)
        r2.font.size = Pt(9.5)

    def text(body):
        p = para()
        r = p.add_run(body)
        r.font.size = Pt(9.5)

    # ---------- header ----------
    contact = list(d.CONTACT) + [d.ADDRESS[lang]]
    if lang == 'de' and getattr(d, 'DOB_DE', None):
        contact.append(d.DOB_DE)

    p = para(space_after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(d.NAME)
    r.bold = True
    r.font.size = Pt(19)
    r.font.color.rgb = ACCENT

    p = para(space_after=2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(c['role'])
    r.font.size = Pt(11)
    r.font.color.rgb = GREY

    p = para(space_after=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(SEP.join(contact))
    r.font.size = Pt(9)

    p = para(space_after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(SEP.join(label for label, _url in d.LINKS))
    r.font.size = Pt(9)
    r.font.color.rgb = ACCENT

    # ---------- body ----------
    render = {'head': head, 'entry': entry, 'bullet': bullet, 'skill': skill, 'text': text}
    for kind, *args in c['sections']:
        render[kind](*(strip_marks(a) for a in args))

    doc.save(out)
    print(out)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('out', nargs='?')
    ap.add_argument('--lang', default='en')
    ap.add_argument('--data', help='CV data file (default: application/cv_data.py)')
    a = ap.parse_args()
    d = load_data(a.data)
    if a.lang not in d.DOCX_CV:
        sys.exit(f"unknown --lang {a.lang!r}: expected one of {', '.join(d.DOCX_CV)}")
    build(a.out or stem(d, a.lang) + '.docx', d, a.lang)
